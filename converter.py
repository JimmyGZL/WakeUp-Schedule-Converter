#!/usr/bin/env python
# -*- coding: utf-8 -*-
r"""
课程表转换工具（单文件合并版）

把原来的 test.py（课表解析：乱格式 xlsx -> 规范 7 列课表）
和 transformer.py（xlsx -> csv，自动填充合并单元格）合并到同一个脚本里。

输入既可以是相对文件名，也可以是绝对路径；省略输入时会在命令行里交互式询问。

用法示例：
    python converter.py test.xlsx
    python converter.py "D:\课表\test.xlsx"
    python converter.py 课表.xlsx --mode raw
    python converter.py test.xlsx -o 我的课表.csv
    python converter.py test.xlsx --list-sheets

详细说明见 readme.md，或运行 python converter.py -h
"""

import argparse
import csv
import os
import re
import sys

import openpyxl
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter

# ==================== 默认配置 ====================
DEFAULT_MODE = 'parse'        # parse=解析课表（原 test.py）；raw=原样导出（原 transformer.py）
SHEET_INDEX = 0               # 默认工作表索引，0 表示第一个表
HEADER_ROWS = 2               # 表头占用行数，第 HEADER_ROWS+1 行对应“第1节”
SKIP_FIRST_COL = True         # 解析模式是否跳过 A 列（A 列通常是“第一节”之类的文字标签）
CSV_HEADER = ['课程名称', '星期', '开始节数', '结束节数', '老师', '地点', '周数']
COL_WIDTHS = [16, 6, 10, 10, 26, 20, 14]
SUFFIX = '_转换结果'           # 默认输出文件名后缀
CN_NUM = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5,
          '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}
# ===================================================


# ==================== 通用工具 ====================

def die(msg):
    """打印错误并退出。"""
    print(f'[错误] {msg}')
    sys.exit(1)


def clean_path(raw):
    """去掉用户可能一起粘贴进来的引号和首尾空白，转成绝对路径。"""
    text = (raw or '').strip().strip('"').strip("'").strip()
    return os.path.abspath(os.path.expanduser(text))


def ask_input():
    """没有通过命令行传入文件时，交互式询问。"""
    try:
        raw = input('请输入课表文件名或绝对路径（例如 test.xlsx）：')
    except EOFError:
        raw = ''
    return raw.strip().strip('"').strip("'")


def resolve_source(raw):
    """校验输入文件，返回绝对路径。支持相对文件名和绝对路径。"""
    if not raw:
        die('没有提供输入文件。可运行：python converter.py 文件名.xlsx')
    path = clean_path(raw)
    if not os.path.exists(path):
        die(f'找不到文件：{path}')
    if os.path.isdir(path):
        die(f'这是一个文件夹，请指定具体文件：{path}')
    low = path.lower()
    if low.endswith('.xls'):
        die('不支持旧版 .xls 格式，请先用 Excel / WPS 另存为 .xlsx 再运行。')
    if not low.endswith(('.xlsx', '.xlsm')):
        print(f'[提示] 文件扩展名不是 .xlsx，仍会尝试用 openpyxl 打开：{path}')
    return path


def load_sheet(src, sheet_index):
    """打开工作簿并返回指定工作表。"""
    try:
        wb = openpyxl.load_workbook(src, data_only=True)
    except Exception as exc:
        die(f'无法打开 {src}：{exc}')
    if sheet_index < 0 or sheet_index >= len(wb.worksheets):
        die(f'工作表索引 {sheet_index} 超出范围（该文件共 {len(wb.worksheets)} 个工作表）。')
    return wb, wb.worksheets[sheet_index]


def write_csv(rows, dst):
    """写入 utf-8-sig 编码的 csv（Excel 直接打开不乱码）。"""
    folder = os.path.dirname(dst)
    if folder and not os.path.isdir(folder):
        die(f'输出目录不存在：{folder}')
    try:
        # encoding='utf-8-sig' 是防止用 Excel 打开 CSV 时中文乱码的关键
        with open(dst, 'w', newline='', encoding='utf-8-sig') as f:
            csv.writer(f).writerows(rows)
    except PermissionError:
        die(f'无法写入 {dst}，该文件可能正被 Excel / WPS 打开，请关闭后重试。')
    except OSError as exc:
        die(f'写入 {dst} 失败：{exc}')
    print(f'[OK] CSV 已生成：{dst}')


def write_xlsx(rows, dst, widths=None):
    """把二维数据写入 xlsx 并做简单美化。"""
    folder = os.path.dirname(dst)
    if folder and not os.path.isdir(folder):
        die(f'输出目录不存在：{folder}')
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '转换结果'
    for row in rows:
        ws.append(row)
    if widths:
        for i, w in enumerate(widths, start=1):
            ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical='center', wrap_text=True)
    try:
        wb.save(dst)
    except PermissionError:
        die(f'无法写入 {dst}，该文件可能正被 Excel / WPS 打开，请关闭后重试。')
    except OSError as exc:
        die(f'写入 {dst} 失败：{exc}')
    print(f'[OK] xlsx 已生成：{dst}')


# ==================== 原 transformer.py：xlsx -> csv ====================

def sheet_to_rows(ws):
    """读取整张表，自动填充合并单元格（取左上角的值），返回二维列表。"""
    merged_data = {}
    for rng in ws.merged_cells.ranges:
        top_left_value = ws.cell(rng.min_row, rng.min_col).value
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                merged_data[(row, col)] = top_left_value

    data = []
    for row in ws.iter_rows():
        row_data = []
        for cell in row:
            val = merged_data.get((cell.row, cell.column), cell.value)
            row_data.append('' if val is None else str(val).strip())
        data.append(row_data)
    return data


# ==================== 原 test.py：课表解析 ====================

def cn_to_int(text):
    """中文数字 / 阿拉伯数字 -> int"""
    text = text.strip()
    if text.isdigit():
        return int(text)
    return CN_NUM.get(text, text)


def convert_location(raw):
    """
    地点格式转换：
        '智慧教室一4-355'  ->  '4号楼智慧教室1'
        '机房二3-201'      ->  '3号楼机房2'
    规则：<教室名><序号><楼号>-<房间号>  ->  <楼号>号楼<教室名><序号>
    """
    if not raw:
        return ''
    raw = raw.strip()

    # 匹配 “4-355” 这样的 “楼号-房间号”
    m = re.search(r'(\d+)\s*[-—－]\s*(\d+)', raw)
    if not m:
        return raw

    building = m.group(1)
    prefix = raw[:m.start()].strip(' （()【】[]')     # '智慧教室一'
    suffix = raw[m.end():].strip(' ）()【】[]')

    # 从 prefix 末尾取出序号
    idx_match = re.search(r'([一二三四五六七八九十\d]+)$', prefix)
    if idx_match:
        room_name = prefix[:idx_match.start()].strip()
        idx = cn_to_int(idx_match.group(1))
    else:
        room_name, idx = prefix, ''

    return f"{building}号楼{room_name}{idx}{suffix}"


def parse_cell_text(text, weekday, start_period, end_period):
    """
    解析单元格内的多行文本，返回 7 个字段的列表。
    约定顺序：第1行=课程名，第2行=老师，第3行=周数，之后是“地点：xxx”“总学时：xxx”
    """
    lines = [l.strip() for l in re.split(r'[\r\n]+', str(text)) if l.strip()]
    if not lines:
        return None

    course = lines[0]
    teachers = lines[1] if len(lines) > 1 else ''
    weeks = lines[2] if len(lines) > 2 else ''

    location_raw = ''
    for line in lines[3:]:
        mm = re.match(r'^地点\s*[:：]\s*(.*)$', line)
        if mm:
            location_raw = mm.group(1).strip()
            break

    return [course, weekday, start_period, end_period,
            teachers, convert_location(location_raw), weeks]


def parse_timetable(ws, header_rows=HEADER_ROWS, skip_first_col=SKIP_FIRST_COL):
    """
    把课表工作表解析成 二维列表（第一行是表头）。
    - 星期：B 列(第2列) -> 星期一(1)，C 列 -> 2 ……
    - 节数：第 header_rows+1 行 -> 第1节，合并单元格跨几行就代表连上几节
    """
    merged_topleft = {}
    for rng in ws.merged_cells.ranges:
        topleft = rng.coord.split(':')[0]
        merged_topleft[topleft] = rng

    rows = [list(CSV_HEADER)]
    for row in ws.iter_rows(min_row=header_rows + 1):
        for cell in row:
            if cell.value is None or not str(cell.value).strip():
                continue
            if skip_first_col and cell.column == 1:
                continue

            weekday = cell.column - 1
            start_period = cell.row - header_rows
            end_period = start_period

            rng = merged_topleft.get(cell.coordinate)
            if rng is not None:
                end_period = start_period + (rng.max_row - rng.min_row)

            rec = parse_cell_text(cell.value, weekday, start_period, end_period)
            if rec:
                rows.append(rec)
    return rows


# ==================== 命令行入口 ====================

def build_parser():
    parser = argparse.ArgumentParser(
        prog='converter.py',
        description='课程表转换工具：xlsx -> 规范课表 csv（合并了原 test.py 与 transformer.py）',
        formatter_class=argparse.RawDescriptionHelpFormatter,
        epilog='示例：\n'
               '  python converter.py test.xlsx\n'
               '  python converter.py "D:\\课表\\test.xlsx" -o 我的课表.csv\n'
               '  python converter.py 课表.xlsx --mode raw\n'
               '  python converter.py test.xlsx --list-sheets\n',
    )
    parser.add_argument('input', nargs='?', default=None,
                        help='输入文件名或绝对路径（省略则交互式询问）')
    parser.add_argument('-m', '--mode', choices=['parse', 'raw'], default=DEFAULT_MODE,
                        help='parse=解析成规范课表（默认）；raw=原样导出并填充合并单元格')
    parser.add_argument('-s', '--sheet', type=int, default=SHEET_INDEX,
                        help=f'工作表索引，0 表示第一个表（默认 {SHEET_INDEX}）')
    parser.add_argument('-o', '--output', default=None,
                        help=f'输出的 csv 路径（默认 输入名{SUFFIX}.csv）')
    parser.add_argument('--xlsx-output', default=None,
                        help=f'parse 模式下输出 xlsx 的路径（默认 输入名{SUFFIX}.xlsx）')
    parser.add_argument('--no-xlsx', action='store_true',
                        help='parse 模式下不额外输出 xlsx，只生成 csv')
    parser.add_argument('--header-rows', type=int, default=HEADER_ROWS,
                        help=f'表头占用行数，第 N+1 行对应第1节（默认 {HEADER_ROWS}）')
    parser.add_argument('--keep-first-col', action='store_true',
                        help='parse 模式不跳过 A 列（默认跳过）')
    parser.add_argument('--list-sheets', action='store_true',
                        help='只列出文件里所有工作表名称后退出')
    return parser


def main(argv=None):
    parser = build_parser()
    args = parser.parse_args(argv)

    raw_input = args.input if args.input else ask_input()
    src = resolve_source(raw_input)
    print(f'输入文件：{src}')

    if args.list_sheets:
        wb = openpyxl.load_workbook(src, data_only=True)
        print(f'共 {len(wb.worksheets)} 个工作表：')
        for i, ws in enumerate(wb.worksheets):
            print(f'  [{i}] {ws.title}')
        return 0

    wb, ws = load_sheet(src, args.sheet)
    print(f'使用工作表 [{args.sheet}]：{ws.title}')

    stem = os.path.splitext(src)[0]
    csv_path = args.output or f'{stem}{SUFFIX}.csv'

    if args.mode == 'raw':
        # ---------- 原 transformer.py 的功能 ----------
        rows = sheet_to_rows(ws)
        print(f'读取到 {len(rows)} 行数据（合并单元格已自动填充）。')
        write_csv(rows, csv_path)
    else:
        # ---------- 原 test.py 的功能 ----------
        rows = parse_timetable(ws, args.header_rows, not args.keep_first_col)
        print(f'解析到 {len(rows) - 1} 条课程记录。')
        write_csv(rows, csv_path)
        if not args.no_xlsx:
            write_xlsx(rows, args.xlsx_output or f'{stem}{SUFFIX}.xlsx', COL_WIDTHS)

    print('全部完成。')
    return 0


if __name__ == '__main__':
    sys.exit(main())
