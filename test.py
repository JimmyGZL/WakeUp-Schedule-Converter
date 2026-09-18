# -*- coding: utf-8 -*-
import re
import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.styles import Alignment
import transformer            # 联动：复用 transformer.py 的 xlsx -> csv 转换
# ==================== 可配置参数 ====================
SRC_FILE = 'test.xlsx'      # 源文件名
DST_FILE = '转换结果.xlsx'    # 输出文件名（中间结果）
CSV_FILE = '转换结果.csv'     # 最终 csv 文件名（由 transformer 生成）
SRC_SHEET_INDEX = 0           # 源表索引，0 表示第一个工作表
HEADER_ROWS = 2               # 表头占用的行数，第 HEADER_ROWS+1 行对应“第1节”
SKIP_FIRST_COL = True         # 是否跳过 A 列（A 列通常是“第一节”之类的文字标签）

CN_NUM = {'一': 1, '二': 2, '三': 3, '四': 4, '五': 5,
          '六': 6, '七': 7, '八': 8, '九': 9, '十': 10}
# ===================================================


def cn_to_int(text):
    """中文数字 / 阿拉伯数字 -> int"""
    text = text.strip()
    if text.isdigit():
        return int(text)
    return CN_NUM.get(text, text)


def convert_location(raw):
    """
    地点格式转换：
        '智慧教室一4-355'  ->  '4号楼智慧教室1'
        '机房二3-201'      ->  '3号楼机房2'
    规则：<教室名><序号><楼号>-<房间号>  ->  <楼号>号楼<教室名><序号>
    """
    if not raw:
        return ''
    raw = raw.strip()

    # 匹配 “4-355” 这样的 “楼号-房间号”
    m = re.search(r'(\d+)\s*[-—－]\s*(\d+)', raw)
    if not m:
        return raw

    building = m.group(1)
    prefix = raw[:m.start()].strip(' （()【】[]')     # '智慧教室一'
    suffix = raw[m.end():].strip(' ）()【】[]')

    # 从 prefix 末尾取出序号
    idx_match = re.search(r'([一二三四五六七八九十\d]+)$', prefix)
    if idx_match:
        room_name = prefix[:idx_match.start()].strip()
        idx = cn_to_int(idx_match.group(1))
    else:
        room_name, idx = prefix, ''

    return f"{building}号楼{room_name}{idx}{suffix}"


def parse_cell_text(text, weekday, start_period, end_period):
    """
    解析单元格内的多行文本，返回 7 个字段的列表。
    约定顺序：第1行=课程名，第2行=老师，第3行=周数，之后是“地点：xxx”“总学时：xxx”
    """
    lines = [l.strip() for l in re.split(r'[\r\n]+', str(text)) if l.strip()]
    if not lines:
        return None

    course = lines[0]
    teachers = lines[1] if len(lines) > 1 else ''
    weeks = lines[2] if len(lines) > 2 else ''

    location_raw = ''
    for line in lines[3:]:
        mm = re.match(r'^地点\s*[:：]\s*(.*)$', line)
        if mm:
            location_raw = mm.group(1).strip()
            break

    return [course, weekday, start_period, end_period,
            teachers, convert_location(location_raw), weeks]


def iter_data_cells(ws):
    """
    遍历所有有值的单元格，自动处理合并单元格（只取左上角，并记录合并跨度）。
    yield: (cell, merged_range_or_None)
    """
    merged_topleft = {}
    for rng in ws.merged_cells.ranges:
        topleft = rng.coord.split(':')[0]
        merged_topleft[topleft] = rng

    for row in ws.iter_rows(min_row=HEADER_ROWS + 1):
        for cell in row:
            if cell.value is None or not str(cell.value).strip():
                continue
            if SKIP_FIRST_COL and cell.column == 1:
                continue
            yield cell, merged_topleft.get(cell.coordinate)


def main():
    wb = openpyxl.load_workbook(SRC_FILE, data_only=True)
    ws = wb.worksheets[SRC_SHEET_INDEX]

    out_wb = openpyxl.Workbook()
    out_ws = out_wb.active
    out_ws.title = '转换结果'
    out_ws.append(['课程名称', '星期', '开始节数', '结束节数', '老师', '地点', '周数'])

    for cell, rng in iter_data_cells(ws):
        # 星期：B 列(第2列) -> 星期一(1)，C 列 -> 2 ……
        weekday = cell.column - 1
        # 节数：第 HEADER_ROWS+1 行 -> 第1节
        start_period = cell.row - HEADER_ROWS
        end_period = start_period
        # 如果是合并单元格，向下合并了几行就代表连上几节
        if rng is not None:
            end_period = start_period + (rng.max_row - rng.min_row)

        rec = parse_cell_text(cell.value, weekday, start_period, end_period)
        if rec:
            out_ws.append(rec)

    # 简单美化
    for i, w in enumerate([16, 6, 10, 10, 26, 20, 14], start=1):
        out_ws.column_dimensions[get_column_letter(i)].width = w
    for row in out_ws.iter_rows():
        for c in row:
            c.alignment = Alignment(vertical='center', wrap_text=True)

    out_wb.save(DST_FILE)
    print(f'[OK] 转换完成，结果已保存到：{DST_FILE}')

    # ==================== 联动 transformer：xlsx -> csv ====================
    transformer.convert_xlsx_to_csv(DST_FILE, CSV_FILE)


if __name__ == '__main__':
    main()